"""
build_report.py - Membuat laporan DOCX otomatis dari model/metrics.json + gambar di static/img
Jalankan SETELAH train.py:  python build_report.py
Isi dulu identitas & tautan di bagian KONFIGURASI.
Hasil: Laporan_Regresi_Linear.docx  (ekspor ke PDF lewat Word: File > Save As > PDF)
"""
import json
import os

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

# ============================ KONFIGURASI (ISI) ============================
NAMA = "Muhammad Fadillah Ramadhan"
NIM = "301240051"
PRODI = "S1 Teknik Informatika"
MATKUL = "Machine Learning"
LINK_GITHUB = "https://github.com/USERNAME/ice-cream-regression"
LINK_APP = "https://NAMA-APP.onrender.com"
LINK_YOUTUBE = "https://youtu.be/XXXXXXXXXXX"
LINK_KAGGLE = "https://www.kaggle.com/datasets/vinicius150987/ice-cream-revenue"
# ===========================================================================

with open(os.path.join("model", "metrics.json"), encoding="utf-8") as f:
    m = json.load(f)
IMG = os.path.join("static", "img")

doc = Document()
sec = doc.sections[0]
sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
sec.left_margin = sec.right_margin = Cm(3)
sec.top_margin = sec.bottom_margin = Cm(2.5)

st = doc.styles["Normal"]
st.font.name = "Times New Roman"
st.font.size = Pt(12)
st.paragraph_format.space_after = Pt(6)
st.paragraph_format.line_spacing = 1.5
for lvl, size in ((1, 14), (2, 12)):
    h = doc.styles[f"Heading {lvl}"]
    h.font.name = "Times New Roman"
    h.font.size = Pt(size)
    h.font.bold = True
    h.font.color.rgb = None


def P(text, bold=False, italic=False, align=None):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold, r.italic = bold, italic
    if align == "c":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align != "l":
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    return p


def B(items):
    for it in items:
        doc.add_paragraph(it, style="List Bullet")


def FIG(file, caption, width=13):
    doc.add_picture(os.path.join(IMG, file), width=Cm(width))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    c = P(caption, italic=True, align="c")
    c.paragraph_format.line_spacing = 1.0


def TABLE(rows, header=True, widths=None):
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = t.cell(i, j)
            cell.text = str(val)
            for p in cell.paragraphs:
                p.paragraph_format.line_spacing = 1.0
                for r in p.runs:
                    r.font.size = Pt(11)
                    r.bold = header and i == 0
            if widths:
                cell.width = Cm(widths[j])
    doc.add_paragraph()


b0, b1 = m["intercept"], m["coef"]

# ------------------------------------------------------------------ COVER
for _ in range(3):
    P("")
P("LAPORAN PROYEK AKHIR MODUL 1", bold=True, align="c")
P("IMPLEMENTASI REGRESI LINEAR SEDERHANA UNTUK PREDIKSI PENDAPATAN TOKO ES KRIM "
  "BERDASARKAN SUHU HARIAN DENGAN APLIKASI WEB FLASK", bold=True, align="c")
P("")
P(f"Mata Kuliah: {MATKUL}", align="c")
P("")
P("Disusun oleh:", align="c")
P(f"{NAMA}", bold=True, align="c")
P(f"NIM {NIM}", align="c")
P(PRODI, align="c")
doc.add_page_break()

# ----------------------------------------------------------------- TAUTAN
doc.add_heading("Tautan Proyek", 1)
TABLE([
    ["Output", "Tautan"],
    ["Repositori GitHub", LINK_GITHUB],
    ["Aplikasi Web (hosting)", LINK_APP],
    ["Video Presentasi YouTube", LINK_YOUTUBE],
    ["Dataset (Kaggle)", LINK_KAGGLE],
], widths=[5, 10])

# ------------------------------------------------------------------ BAB I
doc.add_heading("BAB I PENDAHULUAN", 1)
doc.add_heading("1.1 Latar Belakang", 2)
P("Usaha penjualan es krim termasuk bisnis musiman yang permintaannya sangat dipengaruhi cuaca. "
  "Pada hari yang panas orang cenderung membeli lebih banyak produk dingin, sedangkan pada hari yang "
  "sejuk penjualan menurun. Pemilik usaha yang dapat memperkirakan pendapatan harian dari prakiraan suhu "
  "dapat merencanakan stok bahan baku, jumlah karyawan, dan jam operasional dengan lebih efisien sehingga "
  "mengurangi risiko kehabisan stok maupun pemborosan produk.")
P("Hubungan antara satu variabel masukan (suhu) dan satu variabel keluaran (pendapatan) sesuai untuk "
  "dimodelkan dengan Regresi Linear Sederhana. Proyek ini membangun model tersebut, mengevaluasinya, dan "
  "mengemasnya dalam aplikasi web Flask yang dapat dipakai langsung.")
doc.add_heading("1.2 Rumusan Masalah", 2)
B([
    "Bagaimana pola hubungan antara suhu harian dan pendapatan toko es krim?",
    "Seberapa baik model Regresi Linear Sederhana memprediksi pendapatan dari suhu?",
    "Bagaimana makna intersep dan koefisien model dalam konteks bisnis es krim?",
])
doc.add_heading("1.3 Tujuan", 2)
B([
    "Membangun model Regresi Linear Sederhana untuk memprediksi pendapatan dari suhu.",
    "Mengevaluasi model dengan R-squared (R²) dan RMSE.",
    "Menginterpretasikan parameter model dan mengimplementasikannya pada aplikasi web Flask yang di-deploy.",
])

# ----------------------------------------------------------------- BAB II
doc.add_heading("BAB II JUSTIFIKASI STUDI KASUS DAN DATA", 1)
doc.add_heading("2.1 Justifikasi Pemilihan Studi Kasus", 2)
P("Studi kasus ini dipilih karena (1) hubungan sebab-akibatnya logis: suhu yang lebih tinggi meningkatkan "
  "keinginan konsumen membeli es krim; (2) variabel masukan tunggal, yaitu suhu, mudah diperoleh dari "
  "prakiraan cuaca sehingga model praktis dipakai; (3) keluaran berupa nilai kontinu (pendapatan dalam USD) "
  "sehingga tepat untuk regresi; dan (4) hasilnya langsung bernilai bagi perencanaan operasional usaha kecil.")
doc.add_heading("2.2 Sumber dan Deskripsi Dataset", 2)
P(f"Dataset yang digunakan adalah \"Ice Cream Revenue\" dari Kaggle ({LINK_KAGGLE}). Dataset berisi dua kolom "
  f"numerik: Temperature (suhu dalam °C) sebagai variabel independen (X) dan Revenue (pendapatan dalam USD) "
  f"sebagai variabel dependen (y), dengan {m['n_awal']} baris pengamatan.")
P("Catatan keterbatasan: dataset publik ini sangat bersih dan tidak mencantumkan lokasi, periode, maupun "
  "metode pengumpulan data secara rinci, sehingga hasil model sebaiknya dipandang sebagai ilustrasi "
  "metodologis dan bukan jaminan akurasi untuk toko tertentu.")
doc.add_heading("2.3 Kesesuaian dengan Kriteria Tugas", 2)
B([
    "Memiliki dua kolom numerik dengan hubungan sebab-akibat yang logis (suhu → pendapatan).",
    "Hubungan berpotensi linear, divalidasi secara visual melalui scatterplot (Bab IV).",
    "Berasal dari platform terpercaya dan dapat diunduh publik (Kaggle).",
])

# ---------------------------------------------------------------- BAB III
doc.add_heading("BAB III METODOLOGI", 1)
P("Alur kerja yang digunakan: (1) memuat data, (2) eksplorasi dan pra-pemrosesan, (3) validasi asumsi "
  "linearitas dengan scatterplot, (4) pembagian data latih dan uji 80:20 (random_state = 42), "
  "(5) pelatihan LinearRegression dari Scikit-Learn, (6) evaluasi dengan R² dan RMSE, "
  "(7) penyimpanan model dengan Joblib, dan (8) pembuatan serta deployment aplikasi web Flask dengan Bootstrap.")
P("Model Regresi Linear Sederhana dinyatakan sebagai ŷ = β₀ + β₁x, dengan β₀ intersep dan β₁ koefisien "
  "(kemiringan). Parameter diestimasi dengan metode kuadrat terkecil (ordinary least squares).")
P("Metrik evaluasi: R² mengukur proporsi variasi y yang dijelaskan model; RMSE = √(1/n Σ(yᵢ − ŷᵢ)²) mengukur "
  "rata-rata besar kesalahan prediksi dalam satuan yang sama dengan y (USD).")
TABLE([
    ["Komponen", "Teknologi"],
    ["Bahasa & ML", "Python, Scikit-Learn (LinearRegression)"],
    ["Data & visualisasi", "Pandas, NumPy, Matplotlib, Seaborn"],
    ["Penyimpanan model", "Joblib"],
    ["Web", "Flask, HTML, Bootstrap 5"],
    ["Deployment", "Render / Railway (gunicorn)"],
], widths=[5, 10])

# ----------------------------------------------------------------- BAB IV
doc.add_heading("BAB IV EKSPLORASI DAN PRA-PEMROSESAN DATA", 1)
doc.add_heading("4.1 Statistik Deskriptif", 2)
d = m["desc"]
rows = [["Statistik", "Temperature (°C)", "Revenue (USD)"]]
for k, lab in (("count", "Jumlah"), ("mean", "Rata-rata"), ("std", "Simpangan baku"),
               ("min", "Minimum"), ("25%", "Kuartil 1"), ("50%", "Median"),
               ("75%", "Kuartil 3"), ("max", "Maksimum")):
    rows.append([lab, f"{d['Temperature'][k]:.2f}", f"{d['Revenue'][k]:.2f}"])
TABLE(rows, widths=[5, 5, 5])
doc.add_heading("4.2 Pembersihan Data", 2)
TABLE([
    ["Langkah", "Hasil"],
    ["Jumlah data awal", m["n_awal"]],
    ["Nilai hilang (missing)", m["n_missing"]],
    ["Baris duplikat", m["n_duplikat"]],
    ["Outlier (aturan IQR 1,5×) yang dihapus", m["n_outlier"]],
    ["Pendapatan negatif", m["n_negatif"]],
    ["Jumlah data akhir", m["n_akhir"]],
], widths=[9, 4])
P("Outlier dideteksi dengan aturan IQR pada masing-masing kolom. Pendapatan negatif dihapus karena tidak "
  "masuk akal secara bisnis.")
FIG("distribusi.png", "Gambar 4.1 Distribusi Temperature dan Revenue")
doc.add_heading("4.3 Validasi Asumsi Linearitas", 2)
P(f"Scatterplot pada Gambar 4.2 memperlihatkan tren naik yang jelas. Korelasi Pearson sebesar "
  f"r = {m['korelasi']:.4f} menunjukkan hubungan linear positif yang kuat, sehingga Regresi Linear "
  f"Sederhana layak digunakan.")
FIG("scatter.png", "Gambar 4.2 Scatterplot Temperature vs Revenue")

# ------------------------------------------------------------------ BAB V
doc.add_heading("BAB V HASIL MODEL DAN INTERPRETASI", 1)
doc.add_heading("5.1 Hasil Pelatihan dan Evaluasi", 2)
P(f"Data dibagi menjadi {m['n_train']} data latih dan {m['n_test']} data uji.")
TABLE([
    ["Metrik", "Data Latih", "Data Uji"],
    ["R-squared (R²)", f"{m['r2_train']:.4f}", f"{m['r2_test']:.4f}"],
    ["RMSE (USD)", f"{m['rmse_train']:.2f}", f"{m['rmse_test']:.2f}"],
    ["MAE (USD)", "-", f"{m['mae_test']:.2f}"],
], widths=[5, 5, 5])
gap = abs(m["r2_train"] - m["r2_test"])
P(f"Nilai R² data uji {m['r2_test']:.4f} berarti sekitar {m['r2_test']*100:.1f}% variasi pendapatan dapat "
  f"dijelaskan oleh suhu. RMSE data uji ${m['rmse_test']:.2f} berarti prediksi menyimpang sekitar "
  f"${m['rmse_test']:.2f} dari pendapatan sebenarnya (rata-rata pendapatan ${m['rev_mean']:.2f}). "
  f"Selisih R² latih dan uji hanya {gap:.4f}, sehingga tidak terindikasi overfitting.")
FIG("regresi.png", "Gambar 5.1 Garis regresi pada data latih dan data uji")
doc.add_heading("5.2 Persamaan Model", 2)
P(f"Revenue = {b0:.4f} + {b1:.4f} × Temperature", bold=True, align="c")
doc.add_heading("5.3 Interpretasi Intersep dan Koefisien", 2)
P(f"Koefisien (β₁ = {b1:.4f}): setiap kenaikan suhu 1 °C diperkirakan menaikkan pendapatan harian sekitar "
  f"${b1:.2f}, dengan asumsi faktor lain tetap. Contohnya, kenaikan suhu dari 20 °C ke 30 °C diperkirakan "
  f"menambah pendapatan sekitar ${b1*10:.2f} per hari. Bagi pemilik usaha, ini menjadi dasar kuantitatif "
  f"untuk menaikkan stok saat prakiraan cuaca panas.")
if b0 >= 0:
    P(f"Intersep (β₀ = {b0:.4f}): pada suhu 0 °C model memperkirakan pendapatan dasar sekitar ${b0:.2f} per hari. "
      f"Namun suhu terendah dalam data adalah {m['temp_min']:.1f} °C, sehingga intersep lebih tepat dibaca "
      f"sebagai konstanta penyesuaian garis daripada pendapatan nyata pada suhu 0 °C.")
else:
    P(f"Intersep (β₀ = {b0:.4f}): bernilai negatif sehingga secara harfiah berarti pendapatan negatif pada "
      f"suhu 0 °C, yang tidak masuk akal secara bisnis. Suhu terendah dalam data adalah {m['temp_min']:.1f} °C, "
      f"jadi intersep merupakan hasil ekstrapolasi dan hanya berfungsi sebagai konstanta penyesuaian garis. "
      f"Aplikasi membatasi prediksi minimum pada $0.")
FIG("residual.png", "Gambar 5.2 Analisis residual pada data uji")
P("Pada plot residual, sebaran titik di sekitar garis nol tanpa pola melengkung mendukung asumsi linearitas "
  "dan homoskedastisitas secara visual. Perlu dicek pada hasil Anda bila terlihat pola corong atau lengkung.")
doc.add_heading("5.4 Keterbatasan", 2)
B([
    "Hanya satu variabel (suhu); faktor lain seperti hari libur, lokasi, harga, dan hujan tidak dimodelkan.",
    f"Model hanya andal pada rentang suhu {m['temp_min']:.1f}–{m['temp_max']:.1f} °C; di luar itu prediksi berupa ekstrapolasi.",
    "Korelasi tidak sepenuhnya membuktikan kausalitas.",
    "Dataset publik ini tidak menyebutkan lokasi dan periode sehingga generalisasi ke usaha nyata perlu kehati-hatian.",
])

# ----------------------------------------------------------------- BAB VI
doc.add_heading("BAB VI IMPLEMENTASI APLIKASI WEB DAN DEPLOYMENT", 1)
doc.add_heading("6.1 Penyimpanan Model", 2)
P("Model final disimpan dengan Joblib ke file model/model_regresi_linear.joblib dan dimuat ulang oleh "
  "backend Flask saat aplikasi dijalankan.")
doc.add_heading("6.2 Aplikasi Flask", 2)
P("Antarmuka dibuat dengan Flask dan Bootstrap 5. Formulir menerima satu input, yaitu suhu harian (°C). "
  "Backend memvalidasi input (angka, rentang -30 hingga 60 °C), memanggil model, lalu menampilkan "
  "\"Estimasi Pendapatan Harian\" dalam USD. Jika suhu berada di luar rentang data latih, aplikasi menampilkan "
  "peringatan ekstrapolasi. Tersedia pula endpoint JSON /api/predict?suhu=30.")
doc.add_heading("6.3 Deployment", 2)
P(f"Aplikasi di-deploy ke platform hosting gratis menggunakan gunicorn (Procfile: web: gunicorn app:app). "
  f"Aplikasi dapat diakses di {LINK_APP}. Kode sumber tersedia di {LINK_GITHUB}.")
P("[Sisipkan tangkapan layar aplikasi: halaman awal dan hasil prediksi]", italic=True, align="c")

# ---------------------------------------------------------------- BAB VII
doc.add_heading("BAB VII KESIMPULAN DAN SARAN", 1)
doc.add_heading("7.1 Kesimpulan", 2)
B([
    f"Suhu dan pendapatan memiliki hubungan linear positif yang kuat (r = {m['korelasi']:.4f}).",
    f"Model Regresi Linear Sederhana mencapai R² = {m['r2_test']:.4f} dan RMSE = ${m['rmse_test']:.2f} pada data uji.",
    f"Setiap kenaikan 1 °C diperkirakan menambah pendapatan harian sekitar ${b1:.2f}.",
    "Model berhasil dikemas dalam aplikasi web Flask yang di-deploy dan dapat dipakai untuk estimasi cepat.",
])
doc.add_heading("7.2 Saran", 2)
B([
    "Menambahkan variabel lain (hari libur, hujan, harga) dengan Regresi Linear Berganda.",
    "Menggunakan data nyata dari satu toko dalam periode tertentu agar hasil lebih representatif.",
    "Mencoba model non-linear bila pada data baru pola hubungan melengkung.",
])

# ------------------------------------------------------------ DAFTAR PUSTAKA
doc.add_heading("DAFTAR PUSTAKA", 1)
for ref in [
    f"Kaggle. Ice Cream Revenue Dataset. {LINK_KAGGLE}",
    "Pedregosa, F., et al. (2011). Scikit-learn: Machine Learning in Python. Journal of Machine Learning Research, 12, 2825–2830.",
    "Flask Documentation. https://flask.palletsprojects.com",
    "James, G., Witten, D., Hastie, T., & Tibshirani, R. (2021). An Introduction to Statistical Learning (2nd ed.). Springer.",
]:
    P(ref, align="l")

out = "Laporan_Regresi_Linear.docx"
doc.save(out)
print("Laporan tersimpan:", out)
